"""Builds the downloadable Excel workbook: summary dashboard with charts, ~140 metrics by month and
total, every transaction, daily balances, counterparties, flags and a data dictionary."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .metrics import CH, REGISTRY

INK, MUTED, ACCENT, SOFT, LINE = "1F2A37", "6B7280", "0F766E", "E6F4F1", "D9DEE5"
RISK_COL = {"Low": "15803D", "Medium": "B45309", "High": "B91C1C"}
SEV_FILL = {"High": "FEE2E2", "Medium": "FEF3C7", "Low": "E0F2FE"}

FMT_INR = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
FMT_INR0 = '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'
FMT = {"INR": FMT_INR, "%": '0.00"%"', "count": "#,##0", "days": "#,##0", "ratio": "0.00"}

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=ACCENT)


def _hdr(ws, row, values, col=1, widths=None):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def _month_label(k: str) -> str:
    return datetime.strptime(k, "%Y-%m").strftime("%b %Y")


def build(a: dict, applicant_ref: str, report_id: str, path: Path, brand: str = config.BRAND) -> None:
    wb = Workbook()
    t, monthly, months = a["total"], a["monthly"], a["months"]
    meta = a["meta"]
    labels = [_month_label(m) + ("*" if a["partial"][m] else "") for m in months]

    # ------------------------------------------------------------------ Chart Data (built first, used by charts)
    cd = wb.active
    cd.title = "Chart Data"
    _hdr(cd, 1, ["Month", "Adjusted credits", "Adjusted debits", "Net flow", "Average daily balance", "Closing balance", "EMIs", "Net business inflow"])
    for i, (m, lab) in enumerate(zip(months, labels), start=2):
        mm = monthly[m]
        for j, v in enumerate([lab, mm["credit_adj"], mm["debit_adj"], mm["net_flow"], mm["bal_avg"], mm["bal_close"], mm["emi_debit"], mm["inflow_genuine"]], start=1):
            cell = cd.cell(row=i, column=j, value=v)
            if j > 1:
                cell.number_format = FMT_INR
    n = len(months)
    cd.cell(row=n + 4, column=1, value="Channel")
    cd.cell(row=n + 4, column=2, value="Credits")
    cd.cell(row=n + 4, column=3, value="Debits")
    for i, (k, v) in enumerate(CH.items(), start=n + 5):
        cd.cell(row=i, column=1, value=v)
        cd.cell(row=i, column=2, value=t[f"credit_ch_{k.lower()}"]).number_format = FMT_INR
        cd.cell(row=i, column=3, value=t[f"debit_ch_{k.lower()}"]).number_format = FMT_INR
    cd.column_dimensions["A"].width = 16
    for c in "BCDEFGH":
        cd.column_dimensions[c].width = 18

    # ------------------------------------------------------------------ Summary dashboard
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.column_dimensions["A"].width = 3
    ws["B1"] = f"{brand}  |  Bank Statement Analysis"
    ws["B1"].font = Font(bold=True, size=18, color=INK)
    ws["B2"] = f"Applicant reference: {applicant_ref or '-'}    Report ID: {report_id}    Generated: {datetime.now():%d %b %Y %H:%M}"
    ws["B2"].font = Font(size=9, color=MUTED)
    sc = meta.get("summary_check")
    if sc:
        ws["B3"] = ("Verified: credits, debits and transaction counts match the summary printed on the statement."
                    if sc["verified"] else "Warning: totals computed from the transactions differ from the summary printed on the statement.")
        ws["B3"].font = Font(size=9, bold=True, color=RISK_COL["Low"] if sc["verified"] else RISK_COL["High"])

    info = [("Bank", meta.get("bank") or "-"), ("Account", meta.get("account_masked") or "-"), ("Account holder", meta.get("holder") or "-"),
            ("Account type", meta.get("account_type") or "-"), ("IFSC", meta.get("ifsc") or "-"),
            ("Period", f"{a['rows'][0]['date']:%d %b %Y} to {a['rows'][-1]['date']:%d %b %Y}"), ("Transactions", t["txn_count"]),
            ("Months covered", t["months_covered"])]
    for i, (k, v) in enumerate(info):
        r, c = 4 + i // 4, 2 + (i % 4) * 3
        ws.cell(row=r, column=c, value=k).font = Font(size=9, color=MUTED)
        ws.merge_cells(start_row=r, start_column=c + 1, end_row=r, end_column=c + 2)
        ws.cell(row=r, column=c + 1, value=v).font = Font(bold=True, size=10, color=INK)

    tiles = [
        ("Avg monthly credits (adjusted)", t["avg_monthly_credit"], FMT_INR0, INK),
        ("Avg monthly net inflow", t["avg_monthly_inflow"], FMT_INR0, INK),
        ("Average daily balance", t["bal_avg"], FMT_INR0, INK),
        ("Closing balance", t["bal_close"], FMT_INR0, INK),
        ("EMIs per month", round(t["emi_debit"] / max(t["months_covered"], 0.1), 2), FMT_INR0, INK),
        ("FOIR", t["foir_pct"], '0.0"%"', INK),
        ("Bounced items", t["bounce_count"], "#,##0", RISK_COL["High"] if t["bounce_count"] else INK),
        (f"Risk score ({a['risk']['band']})", t["risk_score"], "0", RISK_COL[a["risk"]["band"]]),
    ]
    for i, (lab, val, fmt, color) in enumerate(tiles):
        r0, c0 = 7 + (i // 4) * 4, 2 + (i % 4) * 3
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 2)
        ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c0 + 2)
        a1 = ws.cell(row=r0, column=c0, value=lab)
        a1.font, a1.alignment = Font(size=9, color=MUTED), Alignment(horizontal="left", indent=1)
        b1 = ws.cell(row=r0 + 1, column=c0, value=val)
        b1.number_format, b1.font = fmt, Font(bold=True, size=20, color=color)
        b1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for rr in range(r0, r0 + 3):
            for cc in range(c0, c0 + 3):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=SOFT)

    # charts
    ch1 = BarChart()
    ch1.type, ch1.title, ch1.height, ch1.width = "col", "Monthly credits vs debits (adjusted)", 8, 16
    ch1.add_data(Reference(cd, min_col=2, max_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
    ch1.set_categories(Reference(cd, min_col=1, min_row=2, max_row=n + 1))
    ch1.y_axis.numFmt, ch1.y_axis.majorGridlines = "#,##0", None
    ch1.y_axis.delete = ch1.x_axis.delete = False
    ws.add_chart(ch1, "B16")
    ch2 = LineChart()
    ch2.title, ch2.height, ch2.width = "Average daily and closing balance", 8, 16
    ch2.add_data(Reference(cd, min_col=5, max_col=6, min_row=1, max_row=n + 1), titles_from_data=True)
    ch2.set_categories(Reference(cd, min_col=1, min_row=2, max_row=n + 1))
    ch2.y_axis.numFmt = "#,##0"
    ch2.y_axis.delete = ch2.x_axis.delete = False
    ws.add_chart(ch2, "G16")
    ch3 = BarChart()
    ch3.type, ch3.title, ch3.height, ch3.width = "bar", "Channel mix (whole period)", 8, 16
    ch3.add_data(Reference(cd, min_col=2, max_col=3, min_row=n + 4, max_row=n + 4 + len(CH)), titles_from_data=True)
    ch3.set_categories(Reference(cd, min_col=1, min_row=n + 5, max_row=n + 4 + len(CH)))
    ch3.y_axis.numFmt = "#,##0"
    ch3.y_axis.delete = ch3.x_axis.delete = False
    ws.add_chart(ch3, "B33")
    ch4 = BarChart()
    ch4.type, ch4.title, ch4.height, ch4.width = "col", "Net cash flow by month", 8, 16
    ch4.add_data(Reference(cd, min_col=4, max_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
    ch4.set_categories(Reference(cd, min_col=1, min_row=2, max_row=n + 1))
    ch4.legend = None
    ch4.y_axis.numFmt = "#,##0"
    ch4.y_axis.delete = ch4.x_axis.delete = False
    ws.add_chart(ch4, "G33")

    r = 50
    ws.cell(row=r, column=2, value="Key flags").font = Font(bold=True, size=12, color=INK)
    r += 1
    if a["flags"]:
        for f in a["flags"]:
            ws.cell(row=r, column=2, value=f["severity"]).fill = PatternFill("solid", fgColor=SEV_FILL[f["severity"]])
            ws.cell(row=r, column=2).font = Font(bold=True, size=9)
            ws.cell(row=r, column=3, value=f["title"]).font = Font(bold=True, size=10)
            ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=12)
            d = ws.cell(row=r, column=6, value=f["detail"])
            d.alignment, d.font = Alignment(wrap_text=True, vertical="top"), Font(size=9, color=MUTED)
            ws.row_dimensions[r].height = 30
            r += 1
    else:
        ws.cell(row=r, column=2, value="No risk flags raised.").font = Font(color=RISK_COL["Low"])
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Adjusted credits remove reversals of the account holder's own debits (for example returned cheques credited back). "
            "Months marked * are partial. All amounts in INR. Figures are computed from the statement PDF supplied and have not been independently verified with the bank.").font = Font(size=8, color=MUTED, italic=True)

    # ------------------------------------------------------------------ Metrics
    mw = wb.create_sheet("Metrics")
    mw.freeze_panes = "E2"
    _hdr(mw, 1, ["Group", "Metric", "Unit", "Total"] + labels)
    row = 2
    group = None
    for key, grp, label, unit, desc, scope in REGISTRY:
        if grp != group:
            c = mw.cell(row=row, column=1, value=grp)
            c.font = Font(bold=True, color=ACCENT)
            for cc in range(1, 5 + n):
                mw.cell(row=row, column=cc).fill = PatternFill("solid", fgColor=SOFT)
            group = grp
            row += 1
        mw.cell(row=row, column=2, value=label)
        mw.cell(row=row, column=3, value=unit).font = Font(color=MUTED, size=9)
        vals = [t.get(key)] + ([monthly[m].get(key) for m in months] if scope == "win" else [None] * n)
        for j, v in enumerate(vals, start=4):
            cell = mw.cell(row=row, column=j, value=v)
            cell.number_format = FMT.get(unit, "General")
            cell.border = BOX
        mw.cell(row=row, column=4).font = Font(bold=True)
        row += 1
    mw.column_dimensions["A"].width = 24
    mw.column_dimensions["B"].width = 44
    mw.column_dimensions["C"].width = 8
    for j in range(4, 5 + n):
        mw.column_dimensions[get_column_letter(j)].width = 16

    # ------------------------------------------------------------------ Transactions
    tw = wb.create_sheet("Transactions")
    tw.freeze_panes = "A2"
    _hdr(tw, 1, ["#", "Date", "Narration", "Debit", "Credit", "Balance", "Channel", "Category", "Counterparty", "Reversal pair", "Bounce", "Recurring"])
    for i, x in enumerate(a["rows"], start=2):
        vals = [i - 1, x["date"], x["narration"], x["amount"] if x["dir"] == "D" else None, x["amount"] if x["dir"] == "C" else None,
                x["balance"], x["channel"], x["category"], x["party"], x["pair_id"], "Yes" if x["bounce"] else None, "Yes" if x["recurring"] else None]
        for j, v in enumerate(vals, start=1):
            c = tw.cell(row=i, column=j, value=v)
            if j == 2:
                c.number_format = "dd-mmm-yyyy"
            elif j in (4, 5, 6):
                c.number_format = FMT_INR
        if x["reversal"]:
            for j in range(1, 13):
                tw.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FEF3C7")
    for col, w in zip("ABCDEFGHIJKL", (6, 12, 60, 14, 14, 15, 11, 18, 26, 9, 8, 9)):
        tw.column_dimensions[col].width = w
    tw.auto_filter.ref = f"A1:L{len(a['rows']) + 1}"

    # ------------------------------------------------------------------ Daily balances
    bw = wb.create_sheet("Daily Balances")
    _hdr(bw, 1, ["Date", "End-of-day balance"])
    for i, (d, b) in enumerate(a["eod"], start=2):
        bw.cell(row=i, column=1, value=d).number_format = "dd-mmm-yyyy"
        bw.cell(row=i, column=2, value=b).number_format = FMT_INR
    bw.column_dimensions["A"].width = 14
    bw.column_dimensions["B"].width = 20

    # ------------------------------------------------------------------ Counterparties
    pw = wb.create_sheet("Counterparties")
    r = 1
    for title, data in (("Top payers (net business inflow)", a["parties_credit"]), ("Top payees", a["parties_debit"])):
        pw.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color=INK)
        r += 1
        _hdr(pw, r, ["Counterparty", "Amount", "Share %", "Transactions", "Active months", "First", "Last"])
        for p in data:
            r += 1
            for j, v in enumerate([p["party"], p["amount"], p["share_pct"], p["count"], p["months"], p["first"], p["last"]], start=1):
                c = pw.cell(row=r, column=j, value=v)
                c.number_format = {2: FMT_INR, 3: '0.00"%"', 6: "dd-mmm-yy", 7: "dd-mmm-yy"}.get(j, "General")
        r += 3
    for col, w in zip("ABCDEFG", (34, 18, 10, 13, 13, 12, 12)):
        pw.column_dimensions[col].width = w
    pw.cell(row=r, column=1, value="Cheque deposits do not name the payer, so they are not attributed to a counterparty.").font = Font(size=9, color=MUTED, italic=True)

    # ------------------------------------------------------------------ Flags, round trips, recurring, categories
    fw = wb.create_sheet("Flags & Patterns")
    r = 1
    fw.cell(row=r, column=1, value=f"Risk score {t['risk_score']} / 100  ({a['risk']['band']})").font = Font(bold=True, size=13, color=RISK_COL[a["risk"]["band"]])
    r += 2
    _hdr(fw, r, ["Severity", "Flag", "Detail", "Points"])
    for f in a["flags"]:
        r += 1
        for j, v in enumerate([f["severity"], f["title"], f["detail"], f["points"]], start=1):
            c = fw.cell(row=r, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        fw.cell(row=r, column=1).fill = PatternFill("solid", fgColor=SEV_FILL[f["severity"]])
    r += 3
    fw.cell(row=r, column=1, value="Round-trip candidates").font = Font(bold=True, size=12)
    r += 1
    _hdr(fw, r, ["Date", "Credit", "Paid out within window", "Counterparty", "Narration"])
    for x in a["round_trips"]:
        r += 1
        for j, v in enumerate([x["date"], x["credit"], x["debited_within_window"], x["party"], x["narration"]], start=1):
            c = fw.cell(row=r, column=j, value=v)
            c.number_format = {1: "dd-mmm-yy", 2: FMT_INR, 3: FMT_INR}.get(j, "General")
    r += 3
    fw.cell(row=r, column=1, value="Recurring debits").font = Font(bold=True, size=12)
    r += 1
    _hdr(fw, r, ["Payee", "Category", "Typical amount", "Occurrences", "Months", "Total paid", "First", "Last"])
    for x in a["recurring"]:
        r += 1
        for j, v in enumerate([x["party"], x["category"], x["median_amount"], x["count"], x["months"], x["total"], x["first"], x["last"]], start=1):
            c = fw.cell(row=r, column=j, value=v)
            c.number_format = {3: FMT_INR, 6: FMT_INR, 7: "dd-mmm-yy", 8: "dd-mmm-yy"}.get(j, "General")
    r += 3
    fw.cell(row=r, column=1, value="Spend and receipts by category").font = Font(bold=True, size=12)
    r += 1
    _hdr(fw, r, ["Category", "Credits (INR)", "Credit count", "Debits (INR)", "Debit count"])
    for x in a["categories"]:
        r += 1
        for j, v in enumerate([x["category"], x["cr_amt"], x["cr_n"], x["dr_amt"], x["dr_n"]], start=1):
            fw.cell(row=r, column=j, value=v).number_format = FMT_INR if j in (2, 4) else "General"
    for col, w in zip("ABCDEFGH", (22, 34, 60, 16, 16, 16, 12, 12)):
        fw.column_dimensions[col].width = w

    # ------------------------------------------------------------------ Data dictionary
    dw = wb.create_sheet("Data Dictionary")
    _hdr(dw, 1, ["Group", "Metric", "Unit", "Columns", "How it is calculated"])
    for i, (key, grp, label, unit, desc, scope) in enumerate(REGISTRY, start=2):
        for j, v in enumerate([grp, label, unit, "Total + each month" if scope == "win" else "Total only", desc], start=1):
            dw.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", (22, 40, 8, 20, 90)):
        dw.column_dimensions[col].width = w
    dw.freeze_panes = "A2"

    wb.move_sheet("Chart Data", offset=len(wb.sheetnames))
    wb.properties.title = f"{brand} statement analysis {report_id}"
    wb.properties.creator = brand
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
